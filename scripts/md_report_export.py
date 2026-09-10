#!/usr/bin/env python3
"""
Export a Farsi (RTL) markdown report to .docx and standalone .html.

Handles the markdown subset used in RTL (e.g. Farsi) reports:
headings (#/##/###), paragraphs with **bold** / *italic* / `code`,
bullet and numbered lists, pipe tables, and --- rules.

Usage: python3 scripts/md_report_export.py report.md
Writes report.docx and report.html next to the source.
"""
import html, re, sys
from pathlib import Path

# ----------------------------------------------------------------- parsing
def parse(md: str):
    blocks, lines, i = [], md.splitlines(), 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.startswith("#"):
            level = len(ln) - len(ln.lstrip("#"))
            blocks.append(("h", level, ln[level:].strip())); i += 1; continue
        if ln.strip() == "---":
            blocks.append(("hr",)); i += 1; continue
        if ln.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows)); continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", ln)
        if m:
            ordered = m.group(2)[0].isdigit(); items = []
            while i < len(lines):
                m2 = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])
                if not m2: break
                items.append(m2.group(3)); i += 1
            blocks.append(("list", ordered, items)); continue
        para = []
        while i < len(lines) and lines[i].strip() and not lines[i].startswith("#") \
                and not lines[i].lstrip().startswith("|") and lines[i].strip() != "---" \
                and not re.match(r"^\s*([-*]|\d+\.)\s+", lines[i]):
            para.append(lines[i].rstrip("  ").rstrip()); i += 1
        blocks.append(("p", " ".join(para)))
    return blocks

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")
def inline_runs(text):
    """-> list of (text, bold, italic, code)"""
    out = []
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**"): out.append((part[2:-2], True, False, False))
        elif part.startswith("`"): out.append((part[1:-1], False, False, True))
        elif part.startswith("*"): out.append((part[1:-1], False, True, False))
        else: out.append((part, False, False, False))
    return out

# ------------------------------------------------------------------- docx
def to_docx(blocks, path: Path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    for s in doc.styles:
        try:
            s.font.name = "Tahoma"
            s.element.rPr.rFonts.set(qn("w:cs"), "Tahoma")
        except Exception:
            pass
    doc.styles["Normal"].font.size = Pt(11)

    def rtl_par(p):
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi"); bidi.set(qn("w:val"), "1"); pPr.append(bidi)

    def add_runs(p, text, size=None, color=None):
        for t, b, it, code in inline_runs(text):
            r = p.add_run(t); r.bold = b; r.italic = it
            if code: r.font.name = "Consolas"; r.font.size = Pt((size or 11) - 1)
            elif size: r.font.size = Pt(size)
            if color: r.font.color.rgb = RGBColor.from_string(color)
            rPr = r._r.get_or_add_rPr()
            rtl = OxmlElement("w:rtl"); rtl.set(qn("w:val"), "1"); rPr.append(rtl)

    for blk in blocks:
        kind = blk[0]
        if kind == "h":
            p = doc.add_heading(level=min(blk[1], 3)); rtl_par(p)
            add_runs(p, blk[2], size={1: 18, 2: 15, 3: 13}[min(blk[1], 3)], color="1F3A5F")
        elif kind == "p":
            p = doc.add_paragraph(); rtl_par(p); add_runs(p, blk[1])
        elif kind == "hr":
            p = doc.add_paragraph(); rtl_par(p)
            pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr")
            btm = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "BBBBBB")):
                btm.set(qn(k), v)
            pbdr.append(btm); pPr.append(pbdr)
        elif kind == "list":
            style = "List Number" if blk[1] else "List Bullet"
            for item in blk[2]:
                p = doc.add_paragraph(style=style); rtl_par(p); add_runs(p, item)
        elif kind == "table":
            rows = blk[1]; ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol); t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            tblPr = t._tbl.tblPr
            bv = OxmlElement("w:bidiVisual"); bv.set(qn("w:val"), "1"); tblPr.append(bv)
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = t.cell(ri, ci); cell.text = ""
                    p = cell.paragraphs[0]; rtl_par(p)
                    txt = row[ci] if ci < len(row) else ""
                    if ri == 0:
                        add_runs(p, f"**{txt}**" if txt else "", size=10)
                        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
                        shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "E8EEF5"); tcPr.append(shd)
                    else:
                        add_runs(p, txt, size=10)
            doc.add_paragraph()
    doc.save(path)

# ------------------------------------------------------------------- html
def to_html(blocks, path: Path, title: str):
    def inl(text):
        out = []
        for t, b, it, code in inline_runs(text):
            e = html.escape(t)
            if code: e = f"<code>{e}</code>"
            if b: e = f"<strong>{e}</strong>"
            if it: e = f"<em>{e}</em>"
            out.append(e)
        return "".join(out)
    body = []
    for blk in blocks:
        k = blk[0]
        if k == "h": body.append(f"<h{blk[1]}>{inl(blk[2])}</h{blk[1]}>")
        elif k == "p": body.append(f"<p>{inl(blk[1])}</p>")
        elif k == "hr": body.append("<hr>")
        elif k == "list":
            tag = "ol" if blk[1] else "ul"
            body.append(f"<{tag}>" + "".join(f"<li>{inl(i)}</li>" for i in blk[2]) + f"</{tag}>")
        elif k == "table":
            rows = blk[1]
            head = "".join(f"<th>{inl(c)}</th>" for c in rows[0])
            rest = "".join("<tr>" + "".join(f"<td>{inl(c)}</td>" for c in r) + "</tr>" for r in rows[1:])
            body.append(f'<div class="tw"><table><thead><tr>{head}</tr></thead><tbody>{rest}</tbody></table></div>')
    css = """
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Vazirmatn:wght@400;500;700&display=swap">
<style>
:root{--bg:#fbfaf7;--fg:#1c2430;--muted:#5b6674;--accent:#1f3a5f;--rule:#d9dee5;--th:#e8eef5;--code:#eef1f5;--codefg:#243447}
@media (prefers-color-scheme: dark){:root:not([data-theme="light"]){--bg:#14181e;--fg:#e6e9ee;--muted:#a3adba;--accent:#8fb4e3;--rule:#2b333d;--th:#1f2732;--code:#1f2732;--codefg:#d6dde7}}
:root[data-theme="dark"]{--bg:#14181e;--fg:#e6e9ee;--muted:#a3adba;--accent:#8fb4e3;--rule:#2b333d;--th:#1f2732;--code:#1f2732;--codefg:#d6dde7}
body{background:var(--bg);color:var(--fg);font-family:Vazirmatn,"Segoe UI",Tahoma,Arial,sans-serif;font-size:16px;line-height:1.9;direction:rtl}
main{max-width:52rem;margin:0 auto;padding:2.5rem 1.25rem 4rem}
h1{font-size:1.9rem;line-height:1.35;color:var(--accent);margin:0 0 1.2rem;text-wrap:balance}
h2{font-size:1.35rem;color:var(--accent);margin:2.4rem 0 .8rem;padding-bottom:.3rem;border-bottom:1px solid var(--rule)}
h3{font-size:1.1rem;margin:1.6rem 0 .5rem}
p{margin:.6rem 0}
ul,ol{padding-right:1.4rem;padding-left:0;margin:.5rem 0}
li{margin:.25rem 0}
hr{border:0;border-top:1px solid var(--rule);margin:2rem 0}
code{font-family:ui-monospace,SFMono-Regular,Menlo,Consolas,monospace;font-size:.85em;background:var(--code);color:var(--codefg);padding:.1em .35em;border-radius:4px;direction:ltr;unicode-bidi:embed}
strong{font-weight:700}
.tw{overflow-x:auto;margin:.8rem 0 1.2rem}
table{border-collapse:collapse;width:100%;font-size:.92rem;font-variant-numeric:tabular-nums}
th,td{border:1px solid var(--rule);padding:.45rem .6rem;text-align:right;vertical-align:top}
th{background:var(--th);font-weight:700;white-space:nowrap}
td:has(> code:only-child){direction:ltr;text-align:left}
</style>"""
    path.write_text(f"<title>{html.escape(title)}</title>{css}\n<main>\n" + "\n".join(body) + "\n</main>\n", encoding="utf-8")

if __name__ == "__main__":
    src = Path(sys.argv[1]); blocks = parse(src.read_text(encoding="utf-8"))
    title = next((b[2] for b in blocks if b[0] == "h" and b[1] == 1), src.stem)
    to_docx(blocks, src.with_suffix(".docx")); to_html(blocks, src.with_suffix(".html"), title)
    print("wrote", src.with_suffix(".docx").name, "and", src.with_suffix(".html").name,
          f"({len(blocks)} blocks, {sum(1 for b in blocks if b[0]=='table')} tables)")
